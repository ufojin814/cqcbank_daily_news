# -*- coding: utf-8 -*-
"""方案二：三张分表"""
import pandas as pd
import numpy as np
import re, os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

TARGET_BANKS = ['重庆农村商业银行','重庆三峡银行','重庆富民银行','重庆银行','成都银行','汉口银行','厦门银行','哈尔滨银行','广东南粤银行','大连银行']
CQB = '重庆银行'
REPORT_DIR = r'C:\Users\Administrator\WorkBuddy\2026-10-16-09-11-37\repo\reports'
DATA1 = r'C:\Users\Administrator\Downloads\银行理财产品查询 (8).xlsx'
DATA2 = r'C:\Users\Administrator\Downloads\银行理财产品查询 (10).xlsx'
BENCH_COL = '业绩比较基准（2026-10-10至2026-10-16）'
FIN_SUBS = ['宁银理财','招银理财','农银理财','杭银理财','华夏理财','北银理财','光大理财']
LOCAL_BANKS = ['重庆三峡银行','重庆农村商业银行','重庆富民银行']
OTHER_BANKS = ['成都银行','汉口银行','厦门银行','哈尔滨银行','广东南粤银行','大连银行']

def safe_float(v):
    try:
        if pd.isna(v): return None
        return float(v)
    except: return None
def get_main_bank(s):
    if pd.isna(s): return ''
    s=str(s)
    # 别名映射：代销机构字段里的写法 → 标准机构名
    alias = {'厦门国际银行':'厦门银行'}
    for k,v in alias.items():
        if k in s: return v
    for b in TARGET_BANKS:
        if b in s: return b
    return '其他'

def get_all_target_banks(s, exclude=None):
    """返回代销机构字段中包含的所有目标银行列表"""
    if pd.isna(s): return []
    s=str(s)
    alias = {'厦门国际银行':'厦门银行'}
    found=[]
    for k,v in alias.items():
        if k in s and v not in found: found.append(v)
    for b in TARGET_BANKS:
        if b in s and b not in found: found.append(b)
    if exclude:
        found=[b for b in found if b not in exclude]
    return found
def parse_benchmark(text):
    if pd.isna(text) or str(text).strip()=='': return None,None
    s=str(text).strip()
    hook=['指数','利率','挂钩','×','＊','*','活期','存款','收益率','中债','国债','信用','利率债','信用债','股票','基金','黄金','原油','商品','汇率','美元','欧元','港元','＋','+','组合','篮子','定存','定期','LPR','shibor','Shibor','SHIBOR','LIBOR']
    if any(k in s for k in hook): return None,None
    nums=re.findall(r'(\d+\.?\d*)\s*%',s)
    if not nums: return None,None
    nums=[float(n) for n in nums]
    if len(nums)==1: return nums[0],nums[0]
    return (min(nums)+max(nums))/2,min(nums)
def get_return(row):
    for col,label in [('近1年年化收益率（%）','近1年'),('近6月年化收益率（%）','近6月'),('近3月年化收益率（%）','近3月'),('近1月年化收益率（%）','近1月')]:
        v=safe_float(row.get(col))
        if v is not None: return v,label
    return None,''
def get_dd(row):
    for col,label in [('近1年最大回撤（%）','近1年'),('近6月最大回撤（%）','近6月'),('近3月最大回撤（%）','近3月'),('近1月最大回撤（%）','近1月')]:
        v=safe_float(row.get(col))
        if v is not None: return v,label
    return None,''

print('读数据...')
df1=pd.read_excel(DATA1,header=1); df2=pd.read_excel(DATA2,header=1)
df=pd.concat([df1,df2],ignore_index=True)
# 过滤产品名称为空的脏记录
df = df.dropna(subset=['产品名称']).reset_index(drop=True)
df['销售机构']=df['代销机构'].apply(get_main_bank)
df=df[df['销售机构'].isin(TARGET_BANKS)].copy()

# 【覆盖修正】表2/表3 候选池用"含任一目标银行"的完整行集：
#   上面按"单一主机构(get_main_bank)"过滤，会因一行同时被多家目标行代销、
#   而第一家是重庆银行时把该行整行剔除，导致其他代销行(汉口/大连/哈尔滨/南粤等)连带丢失。
#   故此处在替换重庆银行之前先留一份完整候选池，供表2/表3使用（表1仍只用按清单匹配的重庆银行行）。
df_share_pool = df.copy()

# 重庆银行替换为新清单匹配到的
cqb_matched = pd.read_pickle(r'C:\Users\Administrator\WorkBuddy\2026-10-16-09-11-37\repo\scripts\_cqb_matched.pkl')
cqb_matched['销售机构']=CQB
df = pd.concat([df[df['销售机构']!=CQB], cqb_matched], ignore_index=True)
print(f'总表: {len(df)}')

df['mid'],df['low'] = zip(*df[BENCH_COL].apply(parse_benchmark))
df['is_num']=df['mid'].notna()

# 表2/表3 候选池同步解析业绩基准（该池在替换重庆银行前留存，需单独计算 mid/low）
df_share_pool['mid'],df_share_pool['low'] = zip(*df_share_pool[BENCH_COL].apply(parse_benchmark))

# 重庆银行各类型基准（最高2只代表的最高中枢/下限）
cqb_ref={}
for st in df['运作子模式'].dropna().unique():
    cn=df[(df['运作子模式']==st)&(df['销售机构']==CQB)&(df['mid'].notna())]
    if len(cn):
        top2=cn.sort_values('mid',ascending=False).head(2)
        cqb_ref[st]=(top2['mid'].max(),top2['low'].max())
print('重庆银行基准:',cqb_ref)

now=datetime.now().strftime('%Y%m%d_%H%M')

# 样式
red_fill=PatternFill('solid',fgColor='fecaca')
red_font=Font(bold=True,color='991b1b',name='微软雅黑')
green_fill=PatternFill('solid',fgColor='dcfce7')
green_font=Font(bold=True,name='微软雅黑')
hdr_font=Font(bold=True,color='FFFFFF',size=12,name='微软雅黑')
hdr_fill=PatternFill('solid',fgColor='1e3a8a')
normal_font=Font(name='微软雅黑',size=11)
center=Alignment(horizontal='center',vertical='center',wrap_text=True)

HEADERS=['代销机构','发行机构','产品名称','业绩基准中枢(%)','业绩基准下限(%)','近1年收益(%)','近1年回撤(%)','产品类型']

def write_sheet(ws, rows_df, title):
    ws.title=title
    for c,h in enumerate(HEADERS,1):
        cell=ws.cell(1,c,h); cell.font=hdr_font; cell.fill=hdr_fill; cell.alignment=center
    r=2
    for _,row in rows_df.iterrows():
        ws.cell(r,1,row['销售机构'])
        ws.cell(r,2,row.get('发行机构',''))
        ws.cell(r,3,row.get('产品名称',''))
        if pd.notna(row.get('mid')):
            ws.cell(r,4,round(row['mid'],2)); ws.cell(r,5,round(row['low'],2))
        else:
            t=str(row.get(BENCH_COL,'无') or '无')
            ws.cell(r,4,t); ws.cell(r,5,t)
        rv,rl=get_return(row); dv,dl=get_dd(row)
        ws.cell(r,6, f'{round(rv,2)}({rl})' if rv and rl!='近1年' else (round(rv,2) if rv else '无'))
        ws.cell(r,7, f'{round(dv,2)}({dl})' if dv and dl!='近1年' else (round(dv,2) if dv else '无'))
        ws.cell(r,8,row.get('运作子模式',''))
        # 标红：同类型双高
        st=row.get('运作子模式','')
        ref=cqb_ref.get(st)
        if ref and pd.notna(row.get('mid')) and row['mid']>ref[0] and row['low']>ref[1]:
            for col in range(1,9): ws.cell(r,col).fill=red_fill; ws.cell(r,col).font=red_font
        elif row['销售机构']==CQB:
            for col in range(1,9): ws.cell(r,col).fill=green_fill; ws.cell(r,col).font=green_font
        else:
            for col in range(1,9): ws.cell(r,col).font=normal_font
        r+=1
    ws.freeze_panes='A2'
    widths=[16,16,40,14,14,14,14,14]
    from openpyxl.utils import get_column_letter
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

# 表1：重庆银行产品
cqb_rows = df[df['销售机构']==CQB].sort_values('mid',ascending=False)
print(f'表1重庆银行: {len(cqb_rows)}只')

# ===== 表2/表3 份额口径（最终版，2026-10-16）=====
# 规则（依用户最终确认口径）：
#   1) 产品登记编码 = 产品级配对键（同一产品的各子份额与总份额共享同一登记编码）。
#   2) 母份额不展示：若某(产品, 银行)已有"子份额"数据，则该对的"总份额"行丢弃。
#      —— 粒度取 (产品, 银行)，而非整个产品：某行仅以母份额记录该银行、其他银行才有子份额时，
#         母份额行必须保留，否则会出现覆盖缺口（实测 90 个银行-产品对属此类）。
#   3) 每行按其"代销机构"字段中包含的目标银行逐家展开（一行可对多家银行）。
#   4) 份额归属可判定性：判断依据 = **该产品所有子份额行的"代销机构"字符串是否完全相同**。
#      - 有差异 → 数据为"份额级"，该银行确实登记在该份额下 → 可按份额精确展示（见 5）。
#      - 全同   → 数据为"产品级聚合"（发行方按产品级上报），无法判定该银行代销哪个份额。
#      实测：1721 个(银行,产品)对中 1453 对(84%)该行只出现在 1 个子份额列表里（天然唯一）；
#            175 对出现在多个份额且列表有差异；93 对出现在多个份额且列表全同（不可定位）。
#   5) 份额级产品的多子份额收敛（对比逻辑，规范 四.5）：中枢+下限双高于重庆银行该类型基准的
#      子份额"全部保留"；其余（未达标）子份额"只保留 1 款"（中枢最高）。
#   6) 产品级聚合（无法定位份额归属）且该组有多个子份额 → **改用母份额(总份额)只展示 1 行**；
#      若该产品无母份额行 → 退化为挑 1 款子份额（优先高于重庆银行者，其次中枢最高）。
#   7) 重庆银行不参与本口径（表1按在售清单原样保留全部份额）。
OTHER_ALL = LOCAL_BANKS + OTHER_BANKS

def _tier_of(row):
    return str(row.get('分级份额类型', '')).strip()

def _banks_of(row):
    return [b for b in get_all_target_banks(row.get('代销机构', ''), exclude=CQB)
            if b in OTHER_ALL]

def _qualifies(row):
    """对比逻辑（规范四.5）：同产品类型内，中枢且下限双高于重庆银行该类型基准。"""
    ref = cqb_ref.get(row.get('运作子模式'))
    if not ref:
        return False
    if pd.isna(row.get('mid')) or pd.isna(row.get('low')):
        return False
    return row['mid'] > ref[0] and row['low'] > ref[1]

# 产品"代销机构"字段粒度判定：该产品所有子份额行的代销机构字符串是否全同
#   True  = 产品级聚合（发行方按产品级上报）-> 无法判定该银行代销哪个份额
#   False = 份额级（各份额列表有差异）-> 可用于定位份额
_agg_key = {}
for _k, _sub in df_share_pool[df_share_pool['分级份额类型'].astype(str).str.strip() == '子份额'].groupby('产品登记编码'):
    _agg_key[str(_k)] = (_sub['代销机构'].astype(str).nunique() == 1)

# 母份额(总份额)行索引：(产品登记编码, 银行) -> 母份额行（供"无法定位份额归属"时展示 1 行）
zong_lookup = {}
for _, row in df_share_pool.iterrows():
    if _tier_of(row) != '总份额':
        continue
    pkey = str(row.get('产品登记编码'))
    for b in _banks_of(row):
        zong_lookup[(pkey, b)] = row

# 第一步：逐行按代销机构展开到目标银行（(产品,银行)已有子份额时丢弃该对的母份额）
zi_pairs = set()   # (产品登记编码, 银行) 已存在子份额
for _, row in df_share_pool.iterrows():
    if _tier_of(row) != '子份额':
        continue
    pkey = str(row.get('产品登记编码'))
    for b in _banks_of(row):
        zi_pairs.add((pkey, b))

step1 = []
for _, row in df_share_pool.iterrows():
    tier = _tier_of(row)
    if tier not in ('子份额', '不分份额', '总份额'):
        continue
    banks = _banks_of(row)
    if not banks:
        continue
    pkey = str(row.get('产品登记编码'))
    for b in banks:
        if tier == '总份额' and (pkey, b) in zi_pairs:
            continue   # 该产品对该行已有精确子份额 -> 丢弃母份额
        r = row.copy()
        r['销售机构'] = b
        step1.append(r)

# 第二步：按 (销售机构, 产品登记编码) 收敛
#   组内有子份额 → 只在子份额中挑；否则用组内全部（不分份额/总份额）。
#   口径（2026-10-16 修订）：高于重庆银行同类型基准的子份额"全部保留"；
#   其余（未达标）子份额"只保留 1 款"（中枢最高）；组内全不达标时同样只留 1 款。
_groups = {}
for r in step1:
    _groups.setdefault((r['销售机构'], str(r.get('产品登记编码'))), []).append(r)

def _score(r):
    return (r['mid'] if pd.notna(r.get('mid')) else -1e9,
            r['low'] if pd.notna(r.get('low')) else -1e9)

final_rows = []
_n_grp_multi = 0
_n_keep_hi = 0
_n_keep_one = 0        # 非达标子份额收敛为 1 款
_n_keep_fallback = 0   # 组内全不达标，留 1 款
_n_use_zong = 0        # 无法定位份额归属 -> 改用母份额 1 行
_n_no_mother = 0       # 无法定位但无母份额 -> 挑 1 款子份额
for (bank, pkey), grp in _groups.items():
    sub = [r for r in grp if _tier_of(r) == '子份额']
    cand = sub if sub else grp          # 有子份额只在子份额中挑；否则用组内全部
    if len(grp) > 1:
        _n_grp_multi += 1
    # 产品级聚合（无法定位份额归属）且该组有多个子份额 -> 用母份额只展示 1 行
    if len(cand) > 1 and _agg_key.get(pkey, False):
        z = zong_lookup.get((pkey, bank))
        if z is not None:
            zr = z.copy()
            zr['销售机构'] = bank
            final_rows.append(zr)
            _n_use_zong += 1
        else:
            final_rows.append(max([r for r in cand if _qualifies(r)] or cand, key=_score))
            _n_no_mother += 1
        continue
    hi = [r for r in cand if _qualifies(r)]
    lo = [r for r in cand if not _qualifies(r)]
    if hi:
        if len(cand) > 1:
            _n_keep_hi += 1
        final_rows.extend(hi)           # 高于重庆银行同类型基准的子份额：全部保留
        if lo:
            _n_keep_one += 1
            final_rows.append(max(lo, key=_score))   # 其余子份额：只保留 1 款
    else:
        if len(cand) > 1:
            _n_keep_fallback += 1
        final_rows.append(max(cand, key=_score))

all_share_rows = pd.DataFrame(final_rows) if final_rows else pd.DataFrame(columns=df.columns)
_tv = all_share_rows['分级份额类型'].astype(str).str.strip()
_n_zi = int((_tv == '子份额').sum()); _n_bu = int((_tv == '不分份额').sum()); _n_zong = int((_tv == '总份额').sum())
print(f'表2表3 份额口径: 展开={len(step1)}行 -> 收敛后={len(all_share_rows)}行 '
      f'(子份额={_n_zi}, 不分份额={_n_bu}, 总份额={_n_zong})')
print(f'  多份额组={_n_grp_multi}; 含达标行组={_n_keep_hi}; 非达标收敛1款={_n_keep_one}; 全低留1款={_n_keep_fallback}')
print(f'  无法定位份额->用母份额1行={_n_use_zong}; 无母份额->挑1款={_n_no_mother}')

# 表2：重庆本土银行（三峡银行、重庆农商行、富民银行）自营及代销，排除重庆银行
loc_rows = all_share_rows[all_share_rows['销售机构'].isin(LOCAL_BANKS)].sort_values(['销售机构','mid'],ascending=[True,False])
print(f'表2本土银行: {len(loc_rows)}只')

# 表3：其他在渝销售城商行（成都、汉口、厦门、哈尔滨、南粤、大连）代销，排除重庆银行
oth_rows = all_share_rows[all_share_rows['销售机构'].isin(OTHER_BANKS)].sort_values(['销售机构','发行机构','mid'],ascending=[True,True,False])
print(f'表3其他城商行: {len(oth_rows)}只')

# 生成Excel
wb=Workbook()
ws1=wb.active; write_sheet(ws1, cqb_rows, '重庆银行产品')
ws2=wb.create_sheet(); write_sheet(ws2, loc_rows, '重庆本土银行产品')
ws3=wb.create_sheet(); write_sheet(ws3, oth_rows, '其他城商行代销产品')

out=os.path.join(REPORT_DIR,f'重庆地区代销理财对比_{now}.xlsx')
wb.save(out)
print(f'\n方案二Excel: {out}')
print(f'表1: {len(cqb_rows)}只 | 表2: {len(loc_rows)}只 | 表3: {len(oth_rows)}只')

# 同时输出产品总表（全量字段）
from openpyxl import Workbook as WB
from openpyxl.utils import get_column_letter as gcl
wb_t=WB(); ws_t=wb_t.active; ws_t.title='产品总表'
ac=list(df.columns)
for c,col in enumerate(ac,1):
    cell=ws_t.cell(1,c,col); cell.font=Font(bold=True,color='FFFFFF',name='微软雅黑'); cell.fill=PatternFill('solid',fgColor='1e3a8a')
for ri,row in enumerate(df.itertuples(index=False),2):
    for ci,val in enumerate(row,1):
        ws_t.cell(ri,ci,val if pd.notna(val) else '')
ws_t.freeze_panes='A2'; ws_t.auto_filter.ref=f'A1:{gcl(len(ac))}{len(df)+1}'
tp=os.path.join(REPORT_DIR,f'重庆地区城商行理财_产品总表_{now}.xlsx')
wb_t.save(tp); print(f'产品总表: {tp}')
